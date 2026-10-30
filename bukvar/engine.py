"""Буквар — OCR двигател.

Зарежда PDF и изображения, подготвя страниците (завъртане, изправяне),
разпознава текста с Tesseract и записва търсим PDF, DOCX и TXT.
"""
from __future__ import annotations

import atexit
import io
import os
import re
import shutil
import statistics
import subprocess
import sys
import tempfile
import threading
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable

import numpy as np
import pypdfium2 as pdfium
import pytesseract
from PIL import Image, ImageOps
from pypdf import PdfReader, PdfWriter, Transformation

from .layout import clean_for_ocr
from .textfix import fix_scripts

RENDER_DPI = 300
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp", ".jp2"}
PDF_EXTS = {".pdf"}
SUPPORTED_EXTS = IMAGE_EXTS | PDF_EXTS

LANG_NAMES = {
    "bul": "Български",
    "eng": "Английски",
    "rus": "Руски",
    "ukr": "Украински",
    "srp": "Сръбски",
    "mkd": "Македонски",
    "deu": "Немски",
    "fra": "Френски",
    "ita": "Италиански",
    "spa": "Испански",
    "tur": "Турски",
    "ell": "Гръцки",
    "ron": "Румънски",
}


# --------------------------------------------------------------------------
# Настройка на Tesseract (вграден в програмата или инсталиран отделно)
# --------------------------------------------------------------------------

def _app_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(getattr(sys, "_MEIPASS", Path(sys.executable).parent))
    return Path(__file__).resolve().parent.parent


def configure_tesseract() -> tuple[str | None, str | None]:
    """Намира tesseract и папката с езиковите данни. Връща (exe, tessdata)."""
    exe_name = "tesseract.exe" if os.name == "nt" else "tesseract"
    candidates = [
        _app_dir() / "tesseract" / exe_name,
        Path(sys.executable).parent / "tesseract" / exe_name,
        Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Tesseract-OCR" / exe_name,
    ]
    exe = next((str(p) for p in candidates if p.is_file()), None) or shutil.which("tesseract")
    if exe:
        pytesseract.pytesseract.tesseract_cmd = exe

    tess_candidates = [
        _app_dir() / "tessdata",
        _app_dir() / "tesseract" / "tessdata",
    ]
    if exe:
        tess_candidates.append(Path(exe).parent / "tessdata")
    tessdata = next((str(p) for p in tess_candidates if (p / "eng.traineddata").is_file()), None)
    if tessdata:
        os.environ["TESSDATA_PREFIX"] = tessdata
    # Tesseract работи по-бързо, когато ние паралелизираме по страници.
    os.environ.setdefault("OMP_THREAD_LIMIT", "1")
    return exe, tessdata


def available_languages() -> list[str]:
    try:
        langs = pytesseract.get_languages(config="")
    except Exception:
        return []
    return [l for l in langs if l not in ("osd", "equ")]


# --------------------------------------------------------------------------
# Модел на документа
# --------------------------------------------------------------------------

@dataclass
class Paragraph:
    text: str
    size: float          # средна височина на реда в пиксели
    heading: bool = False


_TMP_DIR: Path | None = None
_pdfium_lock = threading.Lock()   # pdfium не е безопасен при паралелни нишки


def _tmp_dir() -> Path:
    global _TMP_DIR
    if _TMP_DIR is None:
        _TMP_DIR = Path(tempfile.mkdtemp(prefix="bukvar_"))
        atexit.register(shutil.rmtree, _TMP_DIR, True)
    return _TMP_DIR


def _render_pdf_page(path: Path, index: int, dpi: float) -> Image.Image:
    with _pdfium_lock:
        pdf = pdfium.PdfDocument(str(path))
        try:
            return pdf[index].render(scale=dpi / 72).to_pil().convert("RGB")
        finally:
            pdf.close()


def _open_image_frame(path: Path, index: int) -> Image.Image:
    with Image.open(path) as im:
        im.seek(index)
        return ImageOps.exif_transpose(im.copy()).convert("RGB")


@dataclass
class Page:
    """Една страница. Изображението се зарежда при нужда, не се пази в паметта."""
    source: Path
    index: int                      # номер на страницата в източника (от 0)
    is_pdf: bool = False
    processed_path: Path | None = None  # подготвено (завъртяно/изправено) изображение
    geometry_changed: bool = False
    text: str = ""
    paragraphs: list[Paragraph] = field(default_factory=list)
    text_layer_pdf: bytes | None = None   # невидим текстов слой
    confidence: float | None = None
    text_edited: bool = False
    done: bool = False
    error: str | None = None

    @property
    def label(self) -> str:
        return f"{self.source.name} — стр. {self.index + 1}"

    def original_image(self, dpi: float = RENDER_DPI) -> Image.Image:
        if self.is_pdf:
            return _render_pdf_page(self.source, self.index, dpi)
        img = _open_image_frame(self.source, self.index)
        if dpi != RENDER_DPI:
            img.thumbnail((int(img.width * dpi / RENDER_DPI), int(img.height * dpi / RENDER_DPI)))
        return img

    def image(self) -> Image.Image:
        """Изображението, което се разпознава (подготвеното, ако има такова)."""
        if self.processed_path and self.processed_path.exists():
            with Image.open(self.processed_path) as im:
                return im.convert("RGB")
        return self.original_image()

    def thumbnail(self, size: int = 200) -> Image.Image:
        img = self.original_image(dpi=40) if self.is_pdf else self.image()
        img.thumbnail((size, size))
        return img

    def reset(self) -> None:
        self.processed_path = None
        self.geometry_changed = False
        self.text, self.paragraphs, self.text_layer_pdf = "", [], None
        self.confidence, self.text_edited, self.done, self.error = None, False, False, None


def load_document(path: str | Path) -> list[Page]:
    """Бързо: само преброява страниците, без да ги изобразява."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in PDF_EXTS:
        with _pdfium_lock:
            pdf = pdfium.PdfDocument(str(path))
            n = len(pdf)
            pdf.close()
        return [Page(path, i, is_pdf=True) for i in range(n)]
    if ext in IMAGE_EXTS:
        with Image.open(path) as im:
            n = getattr(im, "n_frames", 1)
        return [Page(path, i) for i in range(n)]
    raise ValueError(f"Неподдържан формат: {path.suffix}")


# --------------------------------------------------------------------------
# Предварителна обработка
# --------------------------------------------------------------------------

def _detect_rotation(img: Image.Image) -> int:
    """Завъртане на 0/90/180/270° по данни от Tesseract OSD."""
    try:
        small = img.copy()
        small.thumbnail((2000, 2000))
        osd = pytesseract.image_to_osd(small, config="--psm 0 -c min_characters_to_try=5")
    except Exception:
        return 0
    m = re.search(r"Rotate:\s*(\d+)", osd)
    conf = re.search(r"Orientation confidence:\s*([\d.]+)", osd)
    if m and conf and float(conf.group(1)) >= 2.0:
        return int(m.group(1)) % 360
    return 0


def _otsu(arr: np.ndarray) -> int:
    hist = np.bincount(arr.ravel(), minlength=256).astype(float)
    total = arr.size
    cum = np.cumsum(hist)
    cum_mean = np.cumsum(hist * np.arange(256))
    w0 = cum / total
    w1 = 1 - w0
    with np.errstate(divide="ignore", invalid="ignore"):
        m0 = cum_mean / cum
        m1 = (cum_mean[-1] - cum_mean) / (total - cum)
        between = w0 * w1 * (m0 - m1) ** 2
    return int(np.nanargmax(between))


def _detect_skew(img: Image.Image) -> float:
    """Ъгъл (в градуси), с който img.rotate() изправя редовете.

    Метод на проекциите: завъртаме умалено копие под различни ъгли и
    избираме този, при който сумите по редове са най-„остри“ —
    тогава редовете текст са строго хоризонтални.
    """
    g = img.convert("L")
    g.thumbnail((1000, 1000))
    arr = np.asarray(g)
    ink = (arr < _otsu(arr)).astype(np.uint8) * 255
    if ink.mean() < 0.5:  # почти празна страница
        return 0.0
    bw = Image.fromarray(ink)

    def score(angle: float) -> float:
        rows = np.asarray(bw.rotate(angle, resample=Image.BILINEAR, fillcolor=0)).sum(axis=1, dtype=np.float64)
        return float(np.sum(np.diff(rows) ** 2))

    def best_of(angles) -> tuple[float, float]:
        scores = [(score(a), a) for a in angles]
        return max(scores)

    base = score(0.0)
    s1, a1 = best_of(np.arange(-10, 10.01, 0.5))
    s2, a2 = best_of(np.arange(a1 - 0.5, a1 + 0.51, 0.05))
    if s2 < base * 1.02:  # няма ясен наклон — не пипаме изображението
        return 0.0
    return round(float(a2), 2)


def preprocess(page: Page, auto_rotate: bool = True, deskew: bool = True) -> Image.Image:
    img = page.original_image()
    changed = False
    if auto_rotate:
        rot = _detect_rotation(img)
        if rot:
            img = img.rotate(-rot, expand=True, fillcolor="white")
            changed = True
    if deskew:
        angle = _detect_skew(img)
        if abs(angle) >= 0.3:
            img = img.rotate(angle, resample=Image.BICUBIC, expand=True, fillcolor="white")
            changed = True
    page.geometry_changed = changed
    if changed:
        page.processed_path = _tmp_dir() / f"{id(page)}_{page.index}.png"
        img.save(page.processed_path, dpi=(RENDER_DPI, RENDER_DPI))
    else:
        page.processed_path = None
    return img


# --------------------------------------------------------------------------
# Разпознаване
# --------------------------------------------------------------------------

@dataclass
class _Word:
    text: str
    left: int
    top: int
    width: int
    height: int
    conf: float

    @property
    def right(self) -> int:
        return self.left + self.width

    @property
    def bottom(self) -> int:
        return self.top + self.height


@dataclass
class _Line:
    block: int
    par: int
    order: float
    words: list[_Word]

    @property
    def left(self) -> int:
        return min(w.left for w in self.words)

    @property
    def right(self) -> int:
        return max(w.right for w in self.words)

    @property
    def top(self) -> int:
        return min(w.top for w in self.words)

    @property
    def bottom(self) -> int:
        return max(w.bottom for w in self.words)

    @property
    def height(self) -> float:
        real = [w.height for w in self.words if w.conf >= 0]
        return float(statistics.median(real or [w.height for w in self.words]))

    def text(self) -> tuple[str, bool]:
        """Текст на реда. Големите празнини (колони на формуляр) стават табулации."""
        words = sorted(self.words, key=lambda w: w.left)
        h = self.height
        out, tabbed = words[0].text, False
        for a, b in zip(words, words[1:]):
            gap = b.left - a.right
            if gap > 2.5 * h:
                out += "\t" + b.text
                tabbed = True
            else:
                out += " " + b.text
        return out, tabbed


_NUMBERED = re.compile(r"^(\d+(\.\d+)*\.?|[IVX]+\.|[а-яa-z]\)|[-•–—*])\s")
_MARKS = ("☐", "☒", "[ръкопис]")
_EXTRA_CONF = 100.0   # псевдодумите (☐ ☒ [ръкопис]) са сигурни


def _collapse_marks(text: str) -> str:
    text = re.sub(r"(\[ръкопис\][\s:.,]*){2,}", "[ръкопис] ", text).rstrip()
    return text


def _lines_from_tsv(data: dict) -> list[_Line]:
    lines: dict[tuple, _Line] = {}
    order = 0
    for i, word in enumerate(data["text"]):
        if not word or not word.strip() or data["conf"][i] < 0:
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        if key not in lines:
            lines[key] = _Line(key[0], key[1], float(order), [])
            order += 1
        lines[key].words.append(_Word(word, data["left"][i], data["top"][i],
                                      data["width"][i], data["height"][i], data["conf"][i]))
    return list(lines.values())


def _drop_noise(lines: list[_Line]) -> list[_Line]:
    """Маха остатъци от ръкопис/печати: думи с много ниска увереност."""
    kept = []
    for ln in lines:
        def noise(w: _Word) -> bool:
            if w.conf >= 30:
                return False
            alnum = sum(ch.isalnum() for ch in w.text)
            # кодове и номера („30.ПП.00.МТ.11/7“) пазим дори при ниска увереност
            if any(ch.isdigit() for ch in w.text) and alnum >= 3:
                return False
            return alnum <= 2 or alnum < len(w.text) / 2
        ws = [w for w in ln.words if not noise(w)]
        # ред само от препинателни знаци и черти („.....“, „| |“) — шум
        if ws and any(ch.isalnum() for w in ws for ch in w.text):
            ln.words = ws
            kept.append(ln)
    return kept


def _merge_rows(lines: list[_Line]) -> list[_Line]:
    """Кратки парчета текст вдясно на същия ред (колона „ДА / НЕ“ на формуляр)
    се присъединяват към реда отляво. Дългите редове (истински колони) не пипаме."""
    def short(ln: _Line) -> bool:
        return len(ln.words) <= 3 and sum(len(w.text) for w in ln.words) <= 24

    changed = True
    while changed:
        changed = False
        for ln in sorted(lines, key=lambda l: -l.left):
            if not short(ln):
                continue
            best = None
            for other in lines:
                if other is ln or other.right > ln.left + 5:
                    continue
                overlap = min(ln.bottom, other.bottom) - max(ln.top, other.top)
                if overlap >= 0.5 * min(ln.bottom - ln.top, other.bottom - other.top):
                    if best is None or other.right > best.right:
                        best = other
            if best is not None and (best.block != ln.block or best.par != ln.par):
                best.words.extend(ln.words)
                lines.remove(ln)
                changed = True
                break
    return lines


def _insert_extras(lines: list[_Line], extras) -> None:
    for e in extras:
        word = _Word(e.text, e.left, e.top, e.width, e.height, _EXTRA_CONF)
        best, best_dist = None, None
        for ln in lines:
            pad = 0.35 * ln.height
            if ln.top - pad <= e.cy <= ln.bottom + pad:
                dist = 0 if ln.left <= e.cx <= ln.right else min(abs(e.cx - ln.left), abs(e.cx - ln.right))
                if best is None or dist < best_dist:
                    best, best_dist = ln, dist
        if best is not None and best_dist < 1200:
            best.words.append(word)
            continue
        # Самостоятелен ред: след най-близкия ред над него
        above = [ln for ln in lines if ln.bottom <= e.cy]
        if above:
            ref = min(above, key=lambda ln: (e.cy - ln.bottom) + 0.2 * abs(e.cx - (ln.left + ln.right) / 2))
            lines.append(_Line(ref.block, -1, ref.order + 0.5 + 0.001 * len(lines), [word]))
        else:
            lines.append(_Line(-1, -1, -1 + 0.001 * len(lines), [word]))


def _build_paragraphs(lines: list[_Line]) -> list[Paragraph]:
    lines = sorted(lines, key=lambda ln: ln.order)
    if not lines:
        return []
    page_text = " ".join(w.text for ln in lines for w in ln.words)
    # Дясната граница на текста в блока — за да познаем кратките (последни) редове
    block_right: dict[int, int] = {}
    for ln in lines:
        block_right[ln.block] = max(block_right.get(ln.block, 0), ln.right)

    paragraphs: list[Paragraph] = []
    cur_text, cur_heights, prev = "", [], None
    prev_tabbed = False
    paragraphs_text_tail = ""

    def flush():
        nonlocal cur_text, cur_heights
        if cur_text.strip():
            paragraphs.append(Paragraph(text=cur_text.strip(), size=statistics.median(cur_heights)))
        cur_text, cur_heights = "", []

    for ln in lines:
        text, tabbed = ln.text()
        letters = sum(ch.isalpha() for ch in text)
        # кратките редове („ДА ☐ HE ☒“, „Р-л c-p YC:“) нямат достатъчно букви —
        # азбуката се определя по цялата страница
        text = fix_scripts(text, None if letters >= 20 else page_text)
        text = _collapse_marks(text)
        h = ln.height
        new_par = prev is None
        if prev is not None:
            ph = prev.height
            prev_text = paragraphs_text_tail
            continues = (abs(ln.left - prev.left) < 2 * ph and 0 <= ln.top - prev.bottom < 0.9 * ph
                         and prev.right >= block_right[prev.block] - 4 * ph
                         and prev_text[-1:] not in ".:;!?" and not _NUMBERED.match(text)
                         and not tabbed and not prev_tabbed)
            # ред, който започва с малка буква, продължава изречението отгоре
            lower_cont = (text[:1].islower() and prev_text[-1:] not in ".:;!?"
                          and not tabbed and not prev_tabbed
                          and not any(m in prev_text for m in _MARKS)
                          and 0 <= ln.top - prev.bottom < 1.8 * ph)
            if lower_cont or (continues and (ln.block != prev.block or ln.par != prev.par)):
                pass
            elif (ln.block != prev.block or ln.par != prev.par
                    or tabbed or prev_tabbed
                    or any(m in prev_text for m in _MARKS)
                    or _NUMBERED.match(text)
                    or ln.top - prev.bottom > 1.2 * max(h, ph)
                    or prev.right < block_right[prev.block] - 4 * ph
                    or abs(h - ph) > 0.35 * max(h, ph)):
                new_par = True
        if new_par:
            flush()
            cur_text = text
        elif cur_text.endswith("-") and text[:1].islower() and len(cur_text) > 1 and cur_text[-2].isalpha():
            cur_text = cur_text[:-1] + text   # пренесена дума
        else:
            cur_text += " " + text
        cur_heights.append(h)
        prev, prev_tabbed = ln, tabbed
        paragraphs_text_tail = text
    flush()

    body = statistics.median(p.size for p in paragraphs) if paragraphs else 0
    for p in paragraphs:
        p.heading = body > 0 and p.size >= body * 1.45 and len(p.text) < 160
    return paragraphs


def _parse_tsv(tsv: str) -> dict:
    cols = ["block_num", "par_num", "line_num", "left", "top", "width", "height", "conf", "text"]
    data: dict[str, list] = {c: [] for c in cols}
    rows = tsv.splitlines()
    if not rows:
        return data
    header = rows[0].split("\t")
    pos = {name: header.index(name) for name in cols}
    for row in rows[1:]:
        parts = row.split("\t")
        if len(parts) < len(header):
            parts += [""] * (len(header) - len(parts))
        for c in cols:
            v = parts[pos[c]]
            data[c].append(v if c == "text" else float(v or -1) if c == "conf" else int(v or 0))
    return data


def ocr_page(page: Page, langs: str, img: Image.Image | None = None,
             clean_forms: bool = True) -> None:
    """Едно минаване на Tesseract дава и текстовия слой (PDF), и позициите (TSV)."""
    img = img or page.image()
    extras = []
    if clean_forms:
        cleaned = clean_for_ocr(img, RENDER_DPI)
        ocr_img, extras = cleaned.image, cleaned.extras
    else:
        ocr_img = img
    with tempfile.TemporaryDirectory(prefix="bukvar_") as tmp:
        src = Path(tmp) / "page.png"
        base = Path(tmp) / "out"
        ocr_img.save(src, dpi=(RENDER_DPI, RENDER_DPI))
        cmd = [pytesseract.pytesseract.tesseract_cmd, str(src), str(base),
               "-l", langs, "--psm", "3", "--dpi", str(RENDER_DPI),
               "-c", "textonly_pdf=1", "-c", "tessedit_create_pdf=1",
               "-c", "tessedit_create_tsv=1", "-c", "tessedit_create_txt=0"]
        flags = 0x08000000 if os.name == "nt" else 0  # CREATE_NO_WINDOW
        res = subprocess.run(cmd, capture_output=True, creationflags=flags)
        if res.returncode != 0:
            raise RuntimeError(res.stderr.decode("utf-8", "replace").strip() or "Tesseract грешка")
        data = _parse_tsv(base.with_suffix(".tsv").read_text(encoding="utf-8"))
        page.text_layer_pdf = base.with_suffix(".pdf").read_bytes()

    lines = _merge_rows(_drop_noise(_lines_from_tsv(data)))
    confs = [w.conf for ln in lines for w in ln.words]
    _insert_extras(lines, extras)
    page.paragraphs = _build_paragraphs(lines)
    page.text = "\n\n".join(p.text for p in page.paragraphs)
    page.confidence = round(sum(confs) / len(confs), 1) if confs else None
    page.text_edited = False
    page.done = True


def process_pages(pages: list[Page], langs: str, auto_rotate: bool = True,
                  deskew: bool = True, workers: int | None = None,
                  clean_forms: bool = True,
                  on_page_done: Callable[[int], None] | None = None,
                  should_stop: Callable[[], bool] | None = None) -> None:
    """Обработва страниците паралелно. on_page_done(i) се вика за всяка готова."""
    workers = workers or max(1, min(len(pages), (os.cpu_count() or 2) - 1))

    def job(i: int) -> int:
        if should_stop and should_stop():
            return i
        try:
            img = preprocess(pages[i], auto_rotate, deskew)
            ocr_page(pages[i], langs, img, clean_forms)
            pages[i].error = None
        except Exception as exc:  # една лоша страница не спира останалите
            pages[i].error = str(exc)
        return i

    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = [pool.submit(job, i) for i in range(len(pages))]
        for f in futures:
            i = f.result()
            if on_page_done and (pages[i].done or pages[i].error):
                on_page_done(i)


# --------------------------------------------------------------------------
# Запис
# --------------------------------------------------------------------------

def _image_page_pdf(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.convert("RGB").save(buf, format="PDF", resolution=RENDER_DPI, quality=80)
    return buf.getvalue()


def export_pdf(pages: Iterable[Page], out_path: str | Path) -> None:
    """Търсим PDF: оригиналната страница + невидим текстов слой."""
    writer = PdfWriter()
    source_cache: dict[Path, PdfReader] = {}
    for page in pages:
        if page.is_pdf and not page.geometry_changed:
            reader = source_cache.setdefault(page.source, PdfReader(str(page.source)))
            writer.add_page(reader.pages[page.index])
        else:
            writer.add_page(PdfReader(io.BytesIO(_image_page_pdf(page.image()))).pages[0])
        target = writer.pages[-1]
        if page.text_layer_pdf:
            layer = PdfReader(io.BytesIO(page.text_layer_pdf)).pages[0]
            tw, th = float(target.mediabox.width), float(target.mediabox.height)
            lw, lh = float(layer.mediabox.width), float(layer.mediabox.height)
            rot = (target.get("/Rotate") or 0) % 360
            if rot in (90, 270):
                tw, th = th, tw
            sx, sy = tw / lw, th / lh
            op = Transformation().scale(sx, sy)
            if rot:
                # Слоят е изчислен върху завъртяната (видимата) страница;
                # връщаме го в координатите на неротирания mediabox.
                W, H = float(target.mediabox.width), float(target.mediabox.height)
                shift = {90: (W, 0), 180: (W, H), 270: (0, H)}[rot]
                op = op.rotate(rot).translate(*shift)
            ox, oy = float(target.mediabox.left), float(target.mediabox.bottom)
            if ox or oy:
                op = op.translate(ox, oy)
            target.merge_transformed_page(layer, op)
    writer.add_metadata({"/Producer": "Буквар OCR", "/Creator": "Буквар OCR"})
    with open(out_path, "wb") as fh:
        writer.write(fh)


def _page_paragraphs(page: Page) -> list[Paragraph]:
    if page.text_edited or not page.paragraphs:
        chunks = [c.strip() for c in re.split(r"\n\s*\n", page.text) if c.strip()]
        return [Paragraph(c.replace("\n", " "), 0) for c in chunks]
    return page.paragraphs


def export_docx(pages: Iterable[Page], out_path: str | Path) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    pages = list(pages)
    for n, page in enumerate(pages):
        for p in _page_paragraphs(page):
            if p.heading:
                h = doc.add_paragraph(style="Heading 2")
                h.add_run(p.text)
            else:
                doc.add_paragraph(p.text)
        if n < len(pages) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.core_properties.author = ""
    doc.save(str(out_path))


def export_txt(pages: Iterable[Page], out_path: str | Path) -> None:
    text = "\n\n\f".join(p.text for p in pages)
    Path(out_path).write_text(text, encoding="utf-8")


EXPORTERS = {".pdf": export_pdf, ".docx": export_docx, ".txt": export_txt}


def export(pages: Iterable[Page], out_path: str | Path) -> None:
    ext = Path(out_path).suffix.lower()
    if ext not in EXPORTERS:
        raise ValueError(f"Неподдържан изходен формат: {ext}")
    EXPORTERS[ext](pages, out_path)
